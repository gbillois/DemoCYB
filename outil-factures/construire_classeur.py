#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Construit le squelette du classeur Outil-Factures.xlsx.

Genere les 9 feuilles, leurs en-tetes, les listes de validation, les mises en
forme conditionnelles et quelques exemples. Le classeur produit ne contient PAS
de macros : l'utilisatrice l'ouvre, l'enregistre en .xlsm puis importe les
modules VBA du dossier vba/ (voir GUIDE-MONTAGE.md).
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

# --- Styles partages -------------------------------------------------------
ENTETE_FILL = PatternFill("solid", fgColor="1F3864")
ENTETE_FONT = Font(bold=True, color="FFFFFF", size=10)
TITRE_FONT = Font(bold=True, size=14, color="1F3864")
SOUS_TITRE_FONT = Font(bold=True, size=11, color="1F3864")
BORDURE = Border(*([Side(style="thin", color="BFBFBF")] * 4))
GRIS = PatternFill("solid", fgColor="F2F2F2")


def style_entetes(ws, entetes, ligne=1):
    """Ecrit une ligne d'en-tetes mise en forme et fige le volet."""
    for col, nom in enumerate(entetes, start=1):
        cel = ws.cell(row=ligne, column=col, value=nom)
        cel.fill = ENTETE_FILL
        cel.font = ENTETE_FONT
        cel.alignment = Alignment(horizontal="center", vertical="center",
                                  wrap_text=True)
        cel.border = BORDURE
        ws.column_dimensions[get_column_letter(col)].width = 20
    ws.row_dimensions[ligne].height = 30
    ws.freeze_panes = ws.cell(row=ligne + 1, column=1)


def liste(valeurs):
    """Cree une validation de type liste deroulante."""
    dv = DataValidation(type="list",
                        formula1='"' + ",".join(valeurs) + '"',
                        allow_blank=True)
    dv.error = "Choisir une valeur dans la liste."
    dv.errorTitle = "Valeur non autorisee"
    return dv


wb = openpyxl.Workbook()

# ==========================================================================
# 1. TABLEAU DE BORD
# ==========================================================================
tdb = wb.active
tdb.title = "Tableau de bord"
tdb.sheet_view.showGridLines = False
tdb["B2"] = "OUTIL DE TRAITEMENT DES FACTURES"
tdb["B2"].font = TITRE_FONT
tdb["B3"] = "Suivi signature - paiement - ajout fournisseur"
tdb["B3"].font = Font(italic=True, color="595959")

tdb["B5"] = "FACTURES - signature"
tdb["B5"].font = SOUS_TITRE_FONT
kpis_factures = [
    ("Recues", 'COUNTIF(\'Suivi Factures\'!N:N,"Recue")'),
    ("Donnees a confirmer", 'COUNTIF(\'Suivi Factures\'!N:N,"Donnees a confirmer")'),
    ("En attente signature", 'COUNTIF(\'Suivi Factures\'!N:N,"En attente signature")'),
    ("Envoyees Docusign", 'COUNTIF(\'Suivi Factures\'!N:N,"Envoyee Docusign")'),
    ("Signees", 'COUNTIF(\'Suivi Factures\'!N:N,"Signee")'),
    ("Classees", 'COUNTIF(\'Suivi Factures\'!N:N,"Classee")'),
]
ligne = 6
for libelle, formule in kpis_factures:
    tdb.cell(row=ligne, column=2, value=libelle).font = Font(size=10)
    c = tdb.cell(row=ligne, column=4, value="=" + formule)
    c.font = Font(bold=True, size=11)
    c.fill = GRIS
    ligne += 1

tdb.cell(row=ligne + 1, column=2, value="PAIEMENTS").font = SOUS_TITRE_FONT
kpis_paie = [
    ("En preparation", 'COUNTIF(\'Suivi Paiements\'!J:J,"En preparation")'),
    ("Envoyes Docusign", 'COUNTIF(\'Suivi Paiements\'!J:J,"Envoye Docusign")'),
    ("Payes", 'COUNTIF(\'Suivi Paiements\'!J:J,"Paye")'),
]
ligne += 2
for libelle, formule in kpis_paie:
    tdb.cell(row=ligne, column=2, value=libelle).font = Font(size=10)
    c = tdb.cell(row=ligne, column=4, value="=" + formule)
    c.font = Font(bold=True, size=11)
    c.fill = GRIS
    ligne += 1

tdb.cell(row=ligne + 1, column=2, value="FOURNISSEURS").font = SOUS_TITRE_FONT
kpis_four = [
    ("Onboarding en cours", 'COUNTIF(\'Onboarding Fournisseurs\'!K:K,"Pieces en attente")'),
    ("A creer", 'COUNTIF(\'Suivi Factures\'!K:K,"A creer")'),
]
ligne += 2
for libelle, formule in kpis_four:
    tdb.cell(row=ligne, column=2, value=libelle).font = Font(size=10)
    c = tdb.cell(row=ligne, column=4, value="=" + formule)
    c.font = Font(bold=True, size=11)
    c.fill = GRIS
    ligne += 1

tdb.cell(row=ligne + 2, column=2,
         value="Actions : Alt+F8 ou les boutons (voir GUIDE-MONTAGE.md) -> "
               "ImporterFactures, DetecterRetoursDocusign, RafraichirTableauBord")
tdb.cell(row=ligne + 2, column=2).font = Font(italic=True, size=9, color="595959")
tdb.column_dimensions["A"].width = 3
tdb.column_dimensions["B"].width = 26
tdb.column_dimensions["C"].width = 3
tdb.column_dimensions["D"].width = 14

# ==========================================================================
# 2. SUIVI FACTURES
# ==========================================================================
sf = wb.create_sheet("Suivi Factures")
ENT_FACTURES = [
    "ID Facture", "Date reception", "Fournisseur", "Numero facture",
    "Date facture", "Montant HT", "Montant TVA", "Montant TTC",
    "Type de depense", "Entite / Destination", "Fournisseur connu ?",
    "Signataires requis", "Fichier PDF", "Statut", "Date envoi Docusign",
    "Date signature", "Emplacement final", "Email source (EntryID)",
    "Remarques",
]
style_entetes(sf, ENT_FACTURES)
STATUTS_FACTURE = ["Recue", "Donnees a confirmer", "En attente signature",
                   "Envoyee Docusign", "Signee", "Classee", "Rejetee"]
TYPES_DEPENSE = ["Societe", "Immeuble", "Travaux"]
dv_statut_f = liste(STATUTS_FACTURE)
dv_type = liste(TYPES_DEPENSE)
dv_connu = liste(["Oui", "Non", "A creer"])
sf.add_data_validation(dv_statut_f)
sf.add_data_validation(dv_type)
sf.add_data_validation(dv_connu)
dv_statut_f.add("N2:N2000")
dv_type.add("I2:I2000")
dv_connu.add("K2:K2000")
# Mise en forme conditionnelle du statut
sf.conditional_formatting.add("N2:N2000",
    CellIsRule(operator="equal", formula=['"Signee"'],
               fill=PatternFill("solid", fgColor="C6EFCE")))
sf.conditional_formatting.add("N2:N2000",
    CellIsRule(operator="equal", formula=['"Classee"'],
               fill=PatternFill("solid", fgColor="C6EFCE")))
sf.conditional_formatting.add("N2:N2000",
    CellIsRule(operator="equal", formula=['"Rejetee"'],
               fill=PatternFill("solid", fgColor="FFC7CE")))
sf.conditional_formatting.add("N2:N2000",
    CellIsRule(operator="equal", formula=['"Donnees a confirmer"'],
               fill=PatternFill("solid", fgColor="FFEB9C")))
for lettre in ("F", "G", "H"):
    for r in range(2, 2001):
        sf[f"{lettre}{r}"].number_format = "# ##0.00 EUR"
sf.column_dimensions["L"].width = 28
sf.column_dimensions["M"].width = 32
sf.column_dimensions["S"].width = 30

# ==========================================================================
# 3. SUIVI PAIEMENTS
# ==========================================================================
sp = wb.create_sheet("Suivi Paiements")
ENT_PAIE = [
    "ID Paiement", "Date creation", "Fournisseur", "Societe emettrice",
    "Factures incluses (IDs)", "Montant total", "Compte debite",
    "Solde connu", "Signataires requis", "Statut", "Date paiement",
    "Remarques",
]
style_entetes(sp, ENT_PAIE)
STATUTS_PAIE = ["En preparation", "Annexes generees", "Envoye Docusign",
                "Signe", "Virement prepare", "Paye"]
dv_statut_p = liste(STATUTS_PAIE)
sp.add_data_validation(dv_statut_p)
dv_statut_p.add("J2:J2000")
for lettre in ("F", "H"):
    for r in range(2, 2001):
        sp[f"{lettre}{r}"].number_format = "# ##0.00 EUR"

# ==========================================================================
# 4. FOURNISSEURS
# ==========================================================================
fo = wb.create_sheet("Fournisseurs")
ENT_FOUR = ["ID Fournisseur", "Nom", "SIRET", "IBAN / RIB", "Email contact",
            "Statut onboarding", "Date creation", "Remarques"]
style_entetes(fo, ENT_FOUR)
dv_statut_four = liste(["Complet", "En cours", "A creer"])
fo.add_data_validation(dv_statut_four)
dv_statut_four.add("F2:F2000")
exemples_four = [
    ["FRN-001", "ACME Travaux SARL", "12345678900012", "FR7630001007941234567890185",
     "compta@acme-travaux.fr", "Complet", "2025-01-15", ""],
    ["FRN-002", "Energie Plus SA", "98765432100021", "FR7630004000031234567890143",
     "facturation@energieplus.fr", "Complet", "2025-02-03", ""],
]
for i, ex in enumerate(exemples_four, start=2):
    for j, val in enumerate(ex, start=1):
        fo.cell(row=i, column=j, value=val)

# ==========================================================================
# 5. ONBOARDING FOURNISSEURS
# ==========================================================================
ob = wb.create_sheet("Onboarding Fournisseurs")
ENT_OB = ["ID Demande", "Date demande", "Nom fournisseur", "Email contact",
          "Kbis recu", "RIB signe recu", "CNI recue", "Callback realise",
          "Trust pair realise", "Annexe 7 signee", "Statut", "Date validation",
          "Remarques"]
style_entetes(ob, ENT_OB)
STATUTS_OB = ["Demande envoyee", "Pieces en attente", "Pieces completes",
              "Valide"]
dv_statut_ob = liste(STATUTS_OB)
dv_oui_non = liste(["Oui", "Non"])
ob.add_data_validation(dv_statut_ob)
ob.add_data_validation(dv_oui_non)
dv_statut_ob.add("K2:K2000")
for lettre in ("E", "F", "G", "H", "I", "J"):
    dv_oui_non.add(f"{lettre}2:{lettre}2000")

# ==========================================================================
# 6. REGLES SIGNATAIRES
# ==========================================================================
rs = wb.create_sheet("Regles Signataires")
ENT_RS = ["Type de depense", "Destination", "Seuil min (EUR)",
          "Seuil max (EUR)", "Signataire 1", "Signataire 2", "Signataire 3",
          "Remarques"]
style_entetes(rs, ENT_RS)
dv_type_rs = liste(TYPES_DEPENSE)
dv_dest_rs = liste(["Societe", "Immeuble"])
rs.add_data_validation(dv_type_rs)
rs.add_data_validation(dv_dest_rs)
dv_type_rs.add("A2:A500")
dv_dest_rs.add("B2:B500")
exemples_regles = [
    ["Societe", "Societe", 0, 5000, "Fund Manager", "", "",
     "Petites depenses societe"],
    ["Societe", "Societe", 5000, 50000, "Fund Manager", "Directeur General",
     "", "Depenses societe moyennes"],
    ["Societe", "Societe", 50000, 9999999, "Fund Manager", "Directeur General",
     "President", "Grosses depenses societe"],
    ["Immeuble", "Immeuble", 0, 10000, "Asset Manager", "", "",
     "Depenses immeuble courantes"],
    ["Immeuble", "Immeuble", 10000, 100000, "Asset Manager", "Fund Manager",
     "", "Depenses immeuble importantes"],
    ["Travaux", "Immeuble", 0, 25000, "Responsable Developpement",
     "Asset Manager", "", "Travaux courants"],
    ["Travaux", "Immeuble", 25000, 9999999, "Responsable Developpement",
     "Asset Manager", "Fund Manager", "Gros travaux / developpement"],
]
for i, ex in enumerate(exemples_regles, start=2):
    for j, val in enumerate(ex, start=1):
        c = rs.cell(row=i, column=j, value=val)
        if j in (3, 4):
            c.number_format = "# ##0 EUR"

# ==========================================================================
# 7. PARAMETRES
# ==========================================================================
pa = wb.create_sheet("Parametres")
style_entetes(pa, ["Parametre", "Valeur"])
pa.column_dimensions["A"].width = 38
pa.column_dimensions["B"].width = 55
PARAMS = [
    ("Dossier Outlook surveille", "Factures a traiter"),
    ("Dossier reseau - A signer", r"\\serveur\Factures\A_signer"),
    ("Dossier reseau - Signees", r"\\serveur\Factures\Signees"),
    ("Dossier reseau - Archives", r"\\serveur\Factures\Archives"),
    ("Dossier local extraction PDF", r"C:\Temp\FacturesPDF"),
    ("Expediteur retour Docusign (filtre)", "dse@docusign.net"),
    ("URL Docusign", "https://account.docusign.com"),
    ("Categorie Outlook - traite", "Traitee par l'outil"),
    ("Extraction OCR niveau 2 (Azure)", "Non"),
    ("Azure DI - Endpoint", ""),
    ("Azure DI - Variable d'env. cle", "AZURE_DI_KEY"),
    ("Utilisateur", ""),
]
for i, (k, v) in enumerate(PARAMS, start=2):
    pa.cell(row=i, column=1, value=k).font = Font(bold=True, size=10)
    pa.cell(row=i, column=2, value=v)
dv_oui_non_p = liste(["Oui", "Non"])
pa.add_data_validation(dv_oui_non_p)
dv_oui_non_p.add("B10")

# ==========================================================================
# 8. MODELES EMAILS
# ==========================================================================
me = wb.create_sheet("Modeles Emails")
style_entetes(me, ["Code modele", "Objet", "Corps"])
me.column_dimensions["A"].width = 26
me.column_dimensions["B"].width = 45
me.column_dimensions["C"].width = 90
MODELES = [
    ("DEMANDE_INFOS_FOURNISSEUR",
     "Creation de compte fournisseur - pieces a fournir",
     "Bonjour,\n\nAfin de vous referencer comme fournisseur, merci de nous "
     "retourner par retour de mail :\n- un extrait Kbis de moins de 3 mois ;\n"
     "- un RIB signe par une personne figurant sur le Kbis ;\n- une copie de "
     "la carte d'identite de cette personne ;\n- vos disponibilites pour un "
     "appel de verification (callback).\n\nUn document Annexe 7 vous sera "
     "ensuite transmis pour signature electronique.\n\nCordialement,"),
    ("RELANCE_FOURNISSEUR",
     "Relance - pieces manquantes pour votre referencement",
     "Bonjour,\n\nNous n'avons pas encore recu l'ensemble des pieces "
     "necessaires a votre referencement. Merci de nous les transmettre dans "
     "les meilleurs delais.\n\nCordialement,"),
    ("RELANCE_SIGNATAIRE",
     "Relance signature Docusign - facture en attente",
     "Bonjour,\n\nUne facture est en attente de votre signature dans "
     "Docusign. Merci de bien vouloir la traiter.\n\nCordialement,"),
]
for i, (code, obj, corps) in enumerate(MODELES, start=2):
    me.cell(row=i, column=1, value=code).font = Font(bold=True, size=10)
    me.cell(row=i, column=2, value=obj)
    c = me.cell(row=i, column=3, value=corps)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    me.row_dimensions[i].height = 120

# ==========================================================================
# 9. JOURNAL
# ==========================================================================
jo = wb.create_sheet("Journal")
style_entetes(jo, ["Horodatage", "Action", "Detail", "Utilisateur"])
jo.column_dimensions["A"].width = 20
jo.column_dimensions["C"].width = 60

wb.save("Outil-Factures.xlsx")
print("Outil-Factures.xlsx genere :", ", ".join(s.title for s in wb.worksheets))
